"""Read user sources as pandas DataFrames. Sources are only read, never modified."""
import atexit
import csv
import json
import shutil
import os
import re
import tempfile
import zipfile

import pandas as pd

CHUNK = 200_000
GEOM_COLUMNS = ["__x", "__y", "__geometry"]


def _encoding(path):
    with open(path, "rb") as f:
        head = f.read(1 << 16)
    for enc in ("utf-8-sig", "gb18030"):
        try:
            head.decode(enc)
            return enc
        except UnicodeDecodeError as e:
            # A multi-byte character cut at the buffer end is still valid UTF-8.
            if enc == "utf-8-sig" and e.start > len(head) - 4:
                return enc
    return "latin-1"


def _delimiter(path, enc):
    if path.lower().endswith(".tsv"):
        return "\t"
    with open(path, encoding=enc, errors="replace") as f:
        head = f.read(1 << 15)
    try:
        return csv.Sniffer().sniff(head, delimiters=",\t;|").delimiter
    except csv.Error:
        return ","


def _crs_from_prj(text, name=""):
    hint = None
    upper = name.upper()
    if "GCJ" in upper:
        hint = "GCJ02"
    elif "BD09" in upper:
        hint = "BD09"
    if not text:
        return hint
    if "GEOGCS" in text and ("WGS_1984" in text or "WGS 84" in text) and "PROJCS" not in text:
        return hint or "WGS84"
    try:
        from pyproj import CRS

        code = CRS.from_wkt(text).to_epsg()
        if code == 4326:
            return hint or "WGS84"
        if code:
            return f"EPSG:{code}"
    except Exception:
        pass
    return hint


def _geometry_record(kind, points, parts=None):
    """Point -> __x/__y; line -> __geometry JSON of its longest part."""
    if not points:
        return {"__x": None, "__y": None, "__geometry": None}
    if kind == "point":
        return {"__x": points[0][0], "__y": points[0][1], "__geometry": None}
    if parts and len(parts) > 1:
        bounds = list(parts) + [len(points)]
        pieces = [points[bounds[i] : bounds[i + 1]] for i in range(len(parts))]
        points = max(pieces, key=len)
    coords = [[round(x, 7), round(y, 7)] for x, y in points]
    return {"__x": None, "__y": None, "__geometry": json.dumps(coords, separators=(",", ":"))}


class FileTable:
    def __init__(self, name, loader, columns, rows, geometry=None, crs_hint=None):
        self.name, self.loader, self.columns, self.rows = name, loader, columns, rows
        self.geometry, self.crs_hint = geometry, crs_hint

    def describe(self):
        return {
            "name": self.name,
            "columns": self.columns,
            "rows": self.rows,
            "geometry": self.geometry,
            "crsHint": self.crs_hint,
        }


def _csv_tables(path, name=None):
    enc = _encoding(path)
    sep = _delimiter(path, enc)
    head = pd.read_csv(path, sep=sep, encoding=enc, nrows=5, dtype=str)
    with open(path, "rb") as f:
        rows = max(sum(1 for _ in f) - 1, 0)

    def load(nrows=None):
        if nrows is not None:
            yield pd.read_csv(path, sep=sep, encoding=enc, nrows=nrows, dtype=str, keep_default_na=False)
            return
        yield from pd.read_csv(path, sep=sep, encoding=enc, chunksize=CHUNK, dtype=str, keep_default_na=False)

    return [FileTable(name or os.path.basename(path), load, list(head.columns), rows)]


def _excel_tables(path):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    out = []
    for ws in wb.worksheets:
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header:
            continue
        columns = [str(c) if c is not None else f"列{i + 1}" for i, c in enumerate(header)]
        sheet = ws.title

        def load(nrows=None, sheet=sheet):
            yield pd.read_excel(path, sheet_name=sheet, nrows=nrows, dtype=str, keep_default_na=False)

        out.append(FileTable(sheet, load, columns, max((ws.max_row or 1) - 1, 0)))
    wb.close()
    return out


def _geojson_tables(path):
    with open(path, encoding=_encoding(path)) as f:
        doc = json.load(f)
    features = doc.get("features") if doc.get("type") == "FeatureCollection" else [doc]
    rows, kinds, props = [], set(), []
    for feat in features:
        geom = feat.get("geometry") or {}
        t = geom.get("type")
        c = geom.get("coordinates")
        if t == "Point":
            g = _geometry_record("point", [c])
            kinds.add("point")
        elif t == "LineString":
            g = _geometry_record("line", c)
            kinds.add("line")
        elif t == "MultiLineString":
            g = _geometry_record("line", max(c, key=len) if c else [])
            kinds.add("line")
        elif t in ("Polygon", "MultiPolygon"):
            ring = c[0] if t == "Polygon" else max((p[0] for p in c), key=len)
            g = _geometry_record("line", ring)
            kinds.add("polygon")
        else:
            g = _geometry_record("point", [])
        p = feat.get("properties") or {}
        for k in p:
            if k not in props:
                props.append(k)
        rows.append({**{k: "" if v is None else str(v) for k, v in p.items()}, **g})
    df = pd.DataFrame(rows, columns=props + GEOM_COLUMNS)
    kind = kinds.pop() if len(kinds) == 1 else None
    hint = _crs_from_prj("", os.path.basename(path)) or "WGS84"

    def load(nrows=None):
        yield df if nrows is None else df.head(nrows)

    return [FileTable(os.path.basename(path), load, props + GEOM_COLUMNS, len(df), kind, hint)]


def _shapefile_tables(path):
    import shapefile

    base = path[:-4]
    prj = open(base + ".prj", encoding="utf-8", errors="replace").read() if os.path.exists(base + ".prj") else ""
    enc = "utf-8"
    if os.path.exists(base + ".cpg"):
        enc = open(base + ".cpg").read().strip() or "utf-8"
    with shapefile.Reader(path, encoding=enc, encodingErrors="replace") as r:
        fields = [f[0] for f in r.fields[1:]]
        count = len(r)
        st = r.shapeType
    kind = "point" if st in (1, 11, 21, 8, 18, 28) else "line" if st in (3, 13, 23) else "polygon" if st in (5, 15, 25) else None

    def load(nrows=None):
        with shapefile.Reader(path, encoding=enc, encodingErrors="replace") as r:
            batch = []
            for i, sr in enumerate(r.iterShapeRecords()):
                if nrows is not None and i >= nrows:
                    break
                rec = {k: "" if v is None else str(v) for k, v in zip(fields, sr.record)}
                pts = [tuple(p[:2]) for p in sr.shape.points]
                rec.update(_geometry_record("point" if kind == "point" else "line", pts, getattr(sr.shape, "parts", None)))
                batch.append(rec)
                if len(batch) >= CHUNK:
                    yield pd.DataFrame(batch, columns=fields + GEOM_COLUMNS)
                    batch = []
            yield pd.DataFrame(batch, columns=fields + GEOM_COLUMNS)

    return [FileTable(os.path.basename(path), load, fields + GEOM_COLUMNS, count, kind, _crs_from_prj(prj, os.path.basename(path)))]


GTFS_TABLES = ["agency", "stops", "routes", "trips", "stop_times", "calendar", "calendar_dates", "shapes", "frequencies"]


def _extract(path):
    tmp = tempfile.mkdtemp(prefix="linggo-src-")
    atexit.register(shutil.rmtree, tmp, True)
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            target = os.path.realpath(os.path.join(tmp, info.filename))
            if not target.startswith(os.path.realpath(tmp) + os.sep):
                raise ValueError("Archive entry escapes the extraction directory")
        z.extractall(tmp)
    return tmp


def _is_gtfs_dir(d):
    return os.path.exists(os.path.join(d, "stops.txt")) and os.path.exists(os.path.join(d, "stop_times.txt"))


def _find(root, predicate):
    for dirpath, _, files in os.walk(root):
        if predicate(dirpath, files):
            return dirpath
    return None


def gtfs_dir(path):
    """Directory holding GTFS txt files for a GTFS zip or folder, else None."""
    if os.path.isdir(path):
        return path if _is_gtfs_dir(path) else None
    if path.lower().endswith(".zip"):
        with zipfile.ZipFile(path) as z:
            names = [os.path.basename(n) for n in z.namelist()]
        if "stops.txt" in names and "stop_times.txt" in names:
            d = _extract(path)
            return _find(d, lambda p, f: "stops.txt" in f)
    return None


def open_tables(source):
    """List the tables of a source. Returns (kind, tables)."""
    t = source.get("type")
    if t == "postgis":
        from .postgis import postgis_tables

        return "postgis", postgis_tables(source)
    path = source.get("path")
    if not isinstance(path, str) or not os.path.isabs(path):
        raise ValueError("请提供本机文件的绝对路径")
    if not os.path.exists(path):
        raise ValueError("文件不存在")
    g = gtfs_dir(path)
    if g:
        tables = []
        for name in GTFS_TABLES:
            p = os.path.join(g, name + ".txt")
            if os.path.exists(p):
                tables += _csv_tables(p, name)
        return "gtfs", tables
    lower = path.lower()
    if lower.endswith((".csv", ".tsv", ".txt")):
        return "table", _csv_tables(path)
    if lower.endswith((".xlsx", ".xlsm")):
        return "table", _excel_tables(path)
    if lower.endswith(".xls"):
        raise ValueError("请另存为 .xlsx 后导入")
    if lower.endswith((".geojson", ".json")):
        return "gis", _geojson_tables(path)
    if lower.endswith(".shp"):
        return "gis", _shapefile_tables(path)
    if lower.endswith(".zip"):
        d = _extract(path)
        shp = []
        for dirpath, _, files in os.walk(d):
            shp += [os.path.join(dirpath, f) for f in files if f.lower().endswith(".shp")]
        if shp:
            return "gis", [t for p in sorted(shp) for t in _shapefile_tables(p)]
        raise ValueError("压缩包中没有 GTFS 或 Shapefile")
    raise ValueError("不支持的文件类型")


def table(source, name):
    kind, tables = open_tables(source)
    for t in tables:
        if t.name == name:
            return kind, t
    if len(tables) == 1 and not name:
        return kind, tables[0]
    raise ValueError(f"找不到数据表 {name}")


def sample_rows(df, n=5):
    return [[("" if v is None else str(v))[:200] for v in row] for row in df.head(n).itertuples(index=False)]
